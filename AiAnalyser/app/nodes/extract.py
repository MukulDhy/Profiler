"""Node 1 — document extraction.

Two backends are supported:

* ``light``   — pypdf + python-docx. Pure Python, no ML models, fast startup.
                This is what the portable/desktop build uses.
* ``docling`` — Docling layout-aware conversion (the development default when
                the package is installed).

Pick one with ``EXTRACTION_MODE=light|docling``. When the variable is unset the
node uses Docling if it can be imported and falls back to the light extractor
otherwise, so a slim install also works without any configuration.
"""

from __future__ import annotations

import importlib.util
import logging
import os
import threading
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

_SUPPORTED_SUFFIXES = {".pdf", ".docx"}
_converter = None
_converter_lock = threading.Lock()


def _docling_available() -> bool:
    return importlib.util.find_spec("docling") is not None


def resolve_mode() -> str:
    """Return the extraction backend that should be used for this process."""
    configured = (os.environ.get("EXTRACTION_MODE") or "").strip().lower()
    if configured in {"light", "docling"}:
        return configured
    return "docling" if _docling_available() else "light"


def _get_converter():
    """Lazily build the Docling converter (model load is expensive)."""
    global _converter
    if _converter is None:
        from docling.document_converter import DocumentConverter

        logger.info("Initialising Docling DocumentConverter")
        _converter = DocumentConverter()
    return _converter


def _validate(path: Path) -> str:
    if not path.exists():
        raise FileNotFoundError(f"Document not found: {path}")

    suffix = path.suffix.lower()
    if suffix not in _SUPPORTED_SUFFIXES:
        raise ValueError(
            f"Unsupported document format '{suffix}'. Supported formats: PDF, DOCX."
        )
    return suffix


def convert_document(file_path: str | Path):
    """Convert a supported document and return the Docling result object."""
    path = Path(file_path)
    _validate(path)

    # Docling's converter is not guaranteed to be thread-safe.
    with _converter_lock:
        return _get_converter().convert(str(path))


# ─────────────────────────────────────────────────────────────────────────────
# Light extraction (portable builds)
# ─────────────────────────────────────────────────────────────────────────────


def _extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001 - one bad page shouldn't abort
            logger.warning("Could not read a page of %s: %s", path.name, exc)
            continue
        if text.strip():
            pages.append(text.strip())
    return "\n\n".join(pages)


def _extract_docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style = (getattr(paragraph.style, "name", "") or "").lower()
        if style.startswith("heading"):
            level = "".join(char for char in style if char.isdigit()) or "1"
            blocks.append(f"{'#' * min(int(level), 6)} {text}")
        elif style.startswith("list"):
            blocks.append(f"- {text}")
        else:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))

    return "\n\n".join(blocks)


def _extract_light(path: Path) -> str:
    suffix = _validate(path)
    if suffix == ".pdf":
        return _extract_pdf_text(path)
    return _extract_docx_text(path)


def extract_markdown(file_path: str | Path) -> str:
    path = Path(file_path)
    mode = resolve_mode()

    markdown = ""
    if mode == "light":
        logger.info("Extracting %s with light extractor", path.name)
        markdown = _extract_light(path)
    else:
        try:
            result = convert_document(path)
            markdown = result.document.export_to_markdown()
        except ImportError:
            logger.warning(
                "Docling is not installed — falling back to the light extractor"
            )
            markdown = _extract_light(path)

    if not markdown or not markdown.strip():
        raise ValueError(
            "No text could be extracted from the document. It may be scanned or empty."
        )

    return markdown


def extract_node(state: dict[str, Any]) -> dict[str, Any]:
    # Text may already be provided (re-analysis path) — skip extraction entirely.
    if state.get("markdown"):
        logger.info("Using %s characters of pre-extracted text", len(state["markdown"]))
        return {}

    markdown = extract_markdown(state["file_path"])
    logger.info("Extracted %s characters from %s", len(markdown), state.get("file_name"))
    return {"markdown": markdown}
